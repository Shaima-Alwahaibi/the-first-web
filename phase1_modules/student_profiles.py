"""Student academic profiles from history available before Spring 2026.

Task 7 copies the latest official semester summary, sums historical semester
earned credits, and classifies each student-course history with the
course-specific passing grade. It does not calculate eligibility or a recommendation.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Mapping

import pandas as pd

from preparation_checks import (
    TARGET_ORDER,
    classify_numeric_value,
    classify_period_frame,
    normalize_course_code,
    write_atomic,
)

TASK7_LOGIC_VERSION = "task7-v2"
CUTOFF_SEMESTER = "2025 Fall"
CUTOFF_PERIOD_ORDER = 2025 * 2 + 1
HOLDOUT_SEMESTER = "2026 Spring"
PROBATION_THRESHOLD = 2.0
ATTEMPTED_CREDITS_SOURCE = "Historical Semester Summary - latest cumulative attempted credits"
EARNED_CREDITS_SOURCE = "Historical Semester Summary - sum of semester earned credits"
EARNED_MATCHED = "Matched"
EARNED_MISMATCH = "Mismatch"
EARNED_NO_OFFICIAL = "Official Historical Value Not Available"
EARNED_INSUFFICIENT = "Insufficient Semester Data"
RISK_RULE_VERSION = "task7-probation-only-v1"
LIST_SEPARATOR = "; "
ACADEMIC_RISK_BASIS = (
    "Research description mapped only from the documented probation rule: "
    "latest SGPA below 2.00 or latest LCGPA below 2.00. This is not an official "
    "UTAS academic-risk policy."
)
CREDITS_NOTE = (
    "Total Credits Attempted is the latest historical Cumulative Attempted Credits. "
    "Total Credits Earned is the sum of Semester Earned Credits through Fall 2025. "
    "Those fields are not the same credit scope. Transcript remark N, defined in the "
    "extracted transcript key as a course repeated to improve the GPA, is included in "
    "Semester Earned Credits and omitted from the cumulative attempted running total. "
    "Earned credits are checked against attempted credits only within one semester scope."
)

# Documented transcript order. A higher rank meets a lower minimum.
# FW sits below F so neither letter satisfies D, D+, C-, or C.
GRADE_RANK: dict[str, int] = {
    "A": 12,
    "A-": 11,
    "B+": 10,
    "B": 9,
    "B-": 8,
    "C+": 7,
    "C": 6,
    "C-": 5,
    "D+": 4,
    "D": 3,
    "F": 1,
    "FW": 0,
}
FAILURE_GRADES = {"F", "FW"}
SPECIALIZATION_LABELS = {
    "software engineering": "Software Engineering",
    "network computing and security": "Network Computing and Security",
    "data science and artificial intelligence": "Data Science and Artificial Intelligence",
    "cyber and information security": "Cyber and Information Security",
}

PROFILE_COLUMNS = [
    "Student Code",
    "Current Level",
    "Current Specialization",
    "Latest Semester",
    "Latest SGPA",
    "Latest LCGPA",
    "Latest CGPA",
    "Total Credits Attempted",
    "Total Credits Earned",
    "Completed Course Count",
    "Completed Courses List",
    "Failed Course Count",
    "Failed Courses List",
    "Withdrawn Course Count",
    "Withdrawn Courses List",
    "Repeated Course Count",
    "Repeated Courses List",
    "Remaining Courses Count",
    "Probation Status",
    "Academic Risk Category",
    "Remaining Courses Status",
    "Is Probation",
    "Academic Risk Basis",
    "Risk Rule Version",
    "Historical Semester Count",
    "Historical Attempt Count",
    "Unique Courses Attempted",
    "Official Credits Attempted",
    "Reconstructed Credits Attempted",
    "Attempted Credits Difference",
    "Official Credits Earned",
    "Reconstructed Credits Earned",
    "Earned Credits Difference",
    "Credits Validation Status",
    "Latest Summary Source",
    "Specialization Source",
    "Profile Validation Status",
    "Manual Review Required",
    "Manual Review Reasons",
    "Cutoff Semester",
    "Task 7 Logic Version",
]

COURSE_STATUS_COLUMNS = [
    "Student Code",
    "Course Code",
    "Course Name",
    "Attempt Count",
    "First Attempt Semester",
    "Latest Attempt Semester",
    "First Grade",
    "Latest Grade",
    "First Result",
    "Latest Result",
    "Latest Grade Point",
    "Minimum Passing Grade",
    "Minimum Passing Grade Point",
    "Ever Passed",
    "Ever Failed",
    "Ever Withdrawn",
    "Eventually Passed",
    "Final Status",
    "Is Completed",
    "Is Current Failed",
    "Is Current Withdrawn",
    "Is Repeated",
    "Completion Semester",
    "Completion Grade",
    "Completion Grade Point",
    "Number of Attempts Before Completion",
    "Passed On First Attempt",
    "Attempt Credit Hours",
    "Completion Credit Hours",
    "Review Notes",
]


def normalize_grade(value: object) -> str | None:
    """Return a compact uppercase grade token. Missing values stay missing."""
    if _missing(value):
        return None
    text = re.sub(r"\s+", "", str(value).strip()).upper()
    return text or None


def meets_course_passing_requirement(
    course_code: object,
    grade: object,
    grade_point: object,
    rule: Mapping[str, object] | None,
) -> bool | None:
    """Return whether ``grade`` meets the course minimum letter.

    The documented letter order decides the result. ``grade_point`` is part of
    the attempt record supplied by the caller and is not compared with 2.0.
    ``None`` means the letter or the rule cannot support a decision.
    A withdrawal does not meet a letter minimum.
    """
    del course_code, grade_point
    if rule is None:
        return None
    minimum = normalize_grade(rule.get("Passing Grade"))
    token = normalize_grade(grade)
    if minimum not in GRADE_RANK:
        return None
    if token == "W":
        return False
    if token not in GRADE_RANK:
        return None
    return GRADE_RANK[token] >= GRADE_RANK[minimum]


def validate_task7_inputs(
    course_history: pd.DataFrame,
    semester_summary: pd.DataFrame,
    passing_rules: pd.DataFrame,
    holdout_courses: pd.DataFrame,
    holdout_semesters: pd.DataFrame,
) -> dict[str, object]:
    """Count the historical boundary and reject holdout rows in feature inputs."""
    _require_columns(
        course_history,
        ["Student Code", "Academic Year", "Term", "Semester", "Course Code", "Grade", "Attempt Number"],
        "course history",
    )
    _require_columns(
        semester_summary,
        ["Student Code", "Academic Year", "Term", "Semester", "SGPA", "LCGPA", "CGPA"],
        "semester summary",
    )
    _require_columns(
        passing_rules,
        ["Course Code", "Passing Grade", "Passing Grade Point"],
        "passing grades",
    )
    history = _reject_holdout(course_history, "course history")
    semesters = _reject_holdout(semester_summary, "semester summary")
    history_students = set(history["Student Code"].dropna())
    semester_students = set(semesters["Student Code"].dropna())
    if history_students != semester_students:
        missing_courses = sorted(semester_students - history_students)
        missing_semesters = sorted(history_students - semester_students)
        raise RuntimeError(
            "Historical course students and semester-summary students differ. "
            f"Missing course rows: {missing_courses[:10]}. "
            f"Missing semester rows: {missing_semesters[:10]}."
        )
    earliest = _boundary_semester(history, minimum=True)
    latest = _boundary_semester(history, minimum=False)
    if latest != CUTOFF_SEMESTER:
        raise RuntimeError(
            f"Latest historical semester is {latest}, expected {CUTOFF_SEMESTER}."
        )
    if int(history["Period Order"].max()) != CUTOFF_PERIOD_ORDER:
        raise RuntimeError("Latest historical period order is not Fall 2025.")
    holdout_course_students = set(holdout_courses["Student Code"].dropna()) if "Student Code" in holdout_courses.columns else set()
    spring_only = sorted(holdout_course_students - history_students)
    rules = _passing_rule_map(passing_rules)
    return {
        "historical_course_rows": int(len(history)),
        "historical_students": int(len(history_students)),
        "historical_semester_rows": int(len(semesters)),
        "earliest_semester": earliest,
        "latest_semester": latest,
        "spring_2026_course_rows_excluded": int(len(holdout_courses)),
        "spring_2026_semester_rows_excluded": int(len(holdout_semesters)),
        "spring_2026_students_excluded_from_features": int(len(holdout_course_students)),
        "spring_2026_students_without_historical_profile": int(len(spring_only)),
        "spring_2026_leakage": 0,
        "passing_rule_codes": int(len(rules)),
        "classified_history": history,
        "classified_semesters": semesters,
        "passing_rule_map": rules,
    }


def get_latest_student_snapshot(semester_summary: pd.DataFrame) -> pd.DataFrame:
    """Return one latest pre-Spring 2026 semester row per student.

    Rows that share a student's latest period are not collapsed into a chosen
    GPA. The snapshot is marked as a conflict and its official values are left
    empty for manual review.
    """
    classified = semester_summary if "Partition" in semester_summary.columns else _reject_holdout(semester_summary, "semester summary")
    if not classified["Partition"].eq("Before Spring 2026").all():
        raise RuntimeError("Latest-snapshot input contains a non-historical semester.")
    work = classified.sort_values(["Student Code", "Period Order", "Source Row ID"], kind="mergesort")
    semester_counts = work.groupby("Student Code", sort=True).size()
    latest_order = work.groupby("Student Code", sort=True)["Period Order"].transform("max")
    latest = work.loc[work["Period Order"].eq(latest_order)].copy()
    conflict_counts = latest.groupby("Student Code", sort=True).size()
    conflicts = set(conflict_counts.loc[conflict_counts.gt(1)].index)
    rows: list[dict[str, object]] = []
    for student_code, group in latest.groupby("Student Code", sort=True):
        conflict = student_code in conflicts
        source = None if conflict else group.iloc[0]
        rows.append({
            "Student Code": student_code,
            "Latest Semester": None if conflict else source["Semester"],
            "Latest Period Order": None if conflict else int(source["Period Order"]),
            "Latest SGPA": None if conflict else source["SGPA"],
            "Latest LCGPA": None if conflict else source["LCGPA"],
            "Latest CGPA": None if conflict else source["CGPA"],
            "Current Level Raw": None if conflict else source["Current Level From Transcript"],
            "Official Credits Attempted": None if conflict else source["Cumulative Attempted Credits"],
            "Historical Semester Count": int(semester_counts.loc[student_code]),
            "Snapshot Conflict": conflict,
            "Latest Summary Source": "historical_data_before_spring_2026.xlsx / Semester Summary",
        })
    snapshot = pd.DataFrame(rows)
    if not snapshot["Student Code"].is_unique:
        raise RuntimeError("Latest snapshot produced more than one row for a student.")
    return snapshot.sort_values("Student Code", kind="mergesort").reset_index(drop=True)


def normalize_course_attempts(course_history: pd.DataFrame) -> pd.DataFrame:
    """Attach period order and normalized course codes. Holdout rows raise."""
    classified = course_history if "Partition" in course_history.columns else _reject_holdout(course_history, "course history")
    if not classified["Partition"].eq("Before Spring 2026").all():
        raise RuntimeError("Course-attempt normalization received a Spring 2026 or unresolved row.")
    work = classified.copy()
    work["Course Code"] = work["Course Code"].map(normalize_course_code)
    if work["Student Code"].isna().any() or work["Course Code"].isna().any():
        raise RuntimeError("A historical attempt is missing a student code or course code.")
    work["Grade Token"] = work["Grade"].map(normalize_grade)
    work["Result Token"] = work["Result"].map(_result_token) if "Result" in work.columns else None
    return work


def summarize_course_attempt_history(
    attempts: pd.DataFrame,
    *,
    course_code: str,
    minimum_grade: str | None,
    minimum_point: float | None,
    student_code: str = "S001",
) -> dict[str, object]:
    """Classify one student-course history without dropping earlier attempts."""
    group = attempts.copy()
    group["Student Code"] = student_code
    group["Course Code"] = course_code
    if "Grade Token" not in group.columns:
        group["Grade Token"] = group["Grade"].map(normalize_grade)
    if "Result Token" not in group.columns:
        group["Result Token"] = group["Result"].map(_result_token) if "Result" in group.columns else None
    rule = None
    if minimum_grade is not None:
        rule = {"Passing Grade": minimum_grade, "Passing Grade Point": minimum_point}
    return _summarize_group(group, rule)


def build_student_course_history(
    student_attempts: pd.DataFrame,
    passing_grade_rules: pd.DataFrame | Mapping[str, Mapping[str, object]],
) -> pd.DataFrame:
    """Return one status row per student and course, in chronological order."""
    rules = (
        passing_grade_rules
        if isinstance(passing_grade_rules, Mapping) and not isinstance(passing_grade_rules, pd.DataFrame)
        else _passing_rule_map(passing_grade_rules)
    )
    normalized = normalize_course_attempts(student_attempts)
    rows = [
        _summarize_group(group, rules.get(str(course_code)))
        for (_, course_code), group in normalized.groupby(["Student Code", "Course Code"], sort=True)
    ]
    status = pd.DataFrame(rows)
    return _order_columns(status, COURSE_STATUS_COLUMNS)


def build_course_status_table(
    course_history: pd.DataFrame,
    passing_grade_rules: pd.DataFrame | Mapping[str, Mapping[str, object]],
) -> pd.DataFrame:
    """Build the normalized course-status table for every historical student."""
    return build_student_course_history(course_history, passing_grade_rules)


def pre_holdout_student_attributes(student_summary: pd.DataFrame) -> pd.DataFrame:
    """Keep specialization and earned credits only when the summary predates Spring 2026.

    Student Summary is a current aggregate. A row labeled Spring 2026 cannot
    supply a feature. Agreed labels are normalized only through the explicit
    capitalization map. Conflicting labels are not chosen.
    """
    required = ["Student Code", "Latest Semester", "Transcript Specialization", "Specialization Path Inferred"]
    _require_columns(student_summary, required, "student summary")
    work = student_summary.copy()
    parsed = work["Latest Semester"].map(_split_semester_text)
    work["Academic Year"] = [item[0] if item else pd.NA for item in parsed]
    work["Term"] = [item[1] if item else pd.NA for item in parsed]
    work["Semester"] = work["Latest Semester"]
    classified = classify_period_frame(work)
    usable = classified.loc[classified["Partition"].eq("Before Spring 2026")].copy()
    rows: list[dict[str, object]] = []
    for student_code, group in usable.groupby("Student Code", sort=True):
        if len(group) != 1:
            rows.append({
                "Student Code": student_code,
                "Current Specialization": pd.NA,
                "Specialization Source": "Student Summary",
                "Official Credits Earned": pd.NA,
                "Attribute Review Reason": "Multiple pre-holdout Student Summary rows",
            })
            continue
        source = group.iloc[0]
        specialization, reason = _agreed_specialization(
            source["Transcript Specialization"],
            source["Specialization Path Inferred"],
        )
        earned = source["Total Credits Earned"] if "Total Credits Earned" in group.columns else pd.NA
        if classify_numeric_value(earned, minimum=0, maximum=None) != "ok":
            earned = pd.NA
            reason = _join_reasons(reason, "Pre-holdout earned credits are missing or invalid")
        rows.append({
            "Student Code": student_code,
            "Current Specialization": specialization,
            "Specialization Source": f"Student Summary latest semester {source['Semester']}",
            "Official Credits Earned": earned,
            "Attribute Review Reason": reason,
        })
    columns = [
        "Student Code",
        "Current Specialization",
        "Specialization Source",
        "Official Credits Earned",
        "Attribute Review Reason",
    ]
    if not rows:
        return pd.DataFrame(columns=columns)
    return pd.DataFrame(rows, columns=columns)


def calculate_total_earned_credits(semester_summary: pd.DataFrame) -> pd.DataFrame:
    """Sum Semester Earned Credits through the historical cutoff.

    A missing, non-numeric, or duplicate semester value leaves that student's
    total blank. Zero is kept when the semester record actually contains zero.
    """
    classified = _historical_semester_frame(semester_summary)
    _require_columns(
        classified,
        ["Student Code", "Semester", "Semester Earned Credits"],
        "semester summary",
    )
    duplicate = classified.duplicated(["Student Code", "Semester"], keep=False).astype(bool)
    numeric = pd.to_numeric(classified["Semester Earned Credits"], errors="coerce")
    missing = classified["Semester Earned Credits"].map(_missing).astype(bool)
    invalid = numeric.isna() & ~missing
    work = pd.DataFrame({
        "Student Code": classified["Student Code"].to_numpy(),
        "Earned Credits": numeric.to_numpy(),
        "Unusable Earned Credit": (duplicate | invalid | missing).to_numpy(),
    })
    rows: list[dict[str, object]] = []
    for student_code, group in work.groupby("Student Code", sort=True):
        if bool(group["Unusable Earned Credit"].any()):
            total = pd.NA
            status = EARNED_INSUFFICIENT
        else:
            total = float(group["Earned Credits"].sum())
            status = "Calculated"
        rows.append({
            "Student Code": student_code,
            "Derived Total Credits Earned": total,
            "Earned Credits Calculation Status": status,
            "Earned Credit Semester Rows": int(len(group)),
        })
    earned = pd.DataFrame(rows)
    if not earned["Student Code"].is_unique:
        raise RuntimeError("Earned-credit totals are not unique by student.")
    return earned


def find_attempted_credit_decreases(semester_summary: pd.DataFrame) -> pd.DataFrame:
    """Return students whose cumulative attempted credits fall across semesters.

    The latest cumulative value is not changed. A decrease is evidence for review.
    """
    classified = _historical_semester_frame(semester_summary)
    _require_columns(classified, ["Student Code", "Cumulative Attempted Credits", "Period Order"], "semester summary")
    sort_columns = ["Student Code", "Period Order"]
    if "Source Row ID" in classified.columns:
        sort_columns.append("Source Row ID")
    work = classified.sort_values(sort_columns, kind="mergesort")
    students: list[object] = []
    for student_code, group in work.groupby("Student Code", sort=True):
        values = pd.to_numeric(group["Cumulative Attempted Credits"], errors="coerce")
        if values.notna().sum() >= 2 and bool(values.diff().lt(0).any()):
            students.append(student_code)
    return pd.DataFrame({"Student Code": students})


def assign_historical_earned_credits(
    profiles: pd.DataFrame,
    semester_summary: pd.DataFrame,
) -> pd.DataFrame:
    """Replace Total Credits Earned with the historical semester sum.

    Official pre-Spring 2026 cumulative earned credits stay in
    ``Official Credits Earned`` and are compared with the sum. They do not
    replace it. Total Credits Attempted is left unchanged.
    """
    if not profiles["Student Code"].is_unique:
        raise RuntimeError('profile["Student Code"].is_unique == True failed.')
    earned = calculate_total_earned_credits(semester_summary)
    decreases = set(find_attempted_credit_decreases(semester_summary)["Student Code"])
    updated = profiles.merge(earned, on="Student Code", how="left", validate="one_to_one")
    missing_calculation = updated["Earned Credits Calculation Status"].isna()
    updated.loc[missing_calculation, "Earned Credits Calculation Status"] = EARNED_INSUFFICIENT
    updated["Total Credits Earned"] = updated["Derived Total Credits Earned"]
    calculated = updated["Earned Credits Calculation Status"].eq("Calculated")
    updated["Total Credits Earned Source"] = pd.NA
    updated.loc[calculated, "Total Credits Earned Source"] = EARNED_CREDITS_SOURCE
    updated["Total Credits Attempted Source"] = ATTEMPTED_CREDITS_SOURCE
    comparison = [
        _compare_earned_credit(official, derived, calculation_status)
        for official, derived, calculation_status in zip(
            updated.get("Official Credits Earned", pd.Series(pd.NA, index=updated.index)),
            updated["Derived Total Credits Earned"],
            updated["Earned Credits Calculation Status"],
        )
    ]
    compared = pd.DataFrame(comparison, columns=["Earned Credits Validation Status", "Official Earned Difference"])
    updated["Earned Credits Validation Status"] = compared["Earned Credits Validation Status"].to_numpy()
    updated["Official Earned Difference"] = compared["Official Earned Difference"].to_numpy()
    scope = _same_scope_credit_flags(semester_summary)
    updated = updated.merge(scope, on="Student Code", how="left", validate="one_to_one")
    notes: list[object] = []
    for student_code, semester_gaps, history_gap in zip(
        updated["Student Code"],
        updated["Semester Earned Above Attempted Rows"],
        updated["Historical Earned Above Attempted Sum"],
    ):
        reasons: list[str] = []
        if pd.notna(semester_gaps) and int(semester_gaps) > 0:
            reasons.append("Semester earned credits exceed semester attempted credits")
        if pd.notna(history_gap) and int(history_gap) > 0:
            reasons.append("Historical earned-credit sum exceeds the historical attempted-credit sum")
        if student_code in decreases:
            reasons.append("Cumulative attempted credits decrease across historical semesters")
        notes.append(" | ".join(reasons) if reasons else pd.NA)
    updated["Credit Review Note"] = notes
    return updated


def _same_scope_credit_flags(semester_summary: pd.DataFrame) -> pd.DataFrame:
    """Compare earned and attempted credits only when both use semester rows.

    Latest cumulative attempted credits are a different field and are not used here.
    """
    classified = _historical_semester_frame(semester_summary)
    _require_columns(
        classified,
        ["Student Code", "Semester Attempted Credits", "Semester Earned Credits"],
        "semester summary",
    )
    work = classified.copy()
    work["Semester Attempted Value"] = pd.to_numeric(work["Semester Attempted Credits"], errors="coerce")
    work["Semester Earned Value"] = pd.to_numeric(work["Semester Earned Credits"], errors="coerce")
    rows: list[dict[str, object]] = []
    for student_code, group in work.groupby("Student Code", sort=True):
        comparable = group["Semester Attempted Value"].notna() & group["Semester Earned Value"].notna()
        semester_gaps = int((group.loc[comparable, "Semester Earned Value"] > group.loc[comparable, "Semester Attempted Value"] + 0.001).sum())
        if bool(comparable.all()) and len(group):
            attempted_sum = float(group["Semester Attempted Value"].sum())
            earned_sum = float(group["Semester Earned Value"].sum())
            history_gap = int(earned_sum > attempted_sum + 0.001)
        else:
            attempted_sum = pd.NA
            history_gap = 0
        rows.append({
            "Student Code": student_code,
            "Historical Semester Attempted Credits Sum": attempted_sum,
            "Semester Earned Above Attempted Rows": semester_gaps,
            "Historical Earned Above Attempted Sum": history_gap,
        })
    return pd.DataFrame(rows)


def build_earned_credit_audit(profiles: pd.DataFrame) -> pd.DataFrame:
    """Return the official-versus-derived earned-credit comparison."""
    columns = [
        "Student Code",
        "Derived Total Credits Earned",
        "Official Historical Total Credits Earned",
        "Difference",
        "Validation Status",
        "Total Credits Attempted",
        "Historical Semester Attempted Credits Sum",
        "Total Credits Attempted Source",
        "Total Credits Earned Source",
        "Credit Review",
    ]
    if "Earned Credits Validation Status" not in profiles.columns:
        return pd.DataFrame(columns=columns)
    audit = pd.DataFrame({
        "Student Code": profiles["Student Code"].astype(str),
        "Derived Total Credits Earned": profiles["Total Credits Earned"],
        "Official Historical Total Credits Earned": profiles["Official Credits Earned"],
        "Difference": profiles["Official Earned Difference"],
        "Validation Status": profiles["Earned Credits Validation Status"],
        "Total Credits Attempted": profiles["Total Credits Attempted"],
        "Historical Semester Attempted Credits Sum": profiles["Historical Semester Attempted Credits Sum"] if "Historical Semester Attempted Credits Sum" in profiles.columns else pd.NA,
        "Total Credits Attempted Source": profiles["Total Credits Attempted Source"],
        "Total Credits Earned Source": profiles["Total Credits Earned Source"],
        "Credit Review": profiles["Credit Review Note"],
    })
    return audit.sort_values("Student Code", kind="mergesort").reset_index(drop=True)


def build_profile_quality_checks(
    profiles: pd.DataFrame,
    *,
    spring_2026_leakage: int,
) -> pd.DataFrame:
    """Return profile integrity checks with PASS, REVIEW, or FAIL."""
    duplicate_codes = int(profiles["Student Code"].duplicated().sum())
    checks = [
        _quality_check("Profile rows", int(len(profiles)), fail=len(profiles) != profiles["Student Code"].nunique()),
        _quality_check("Unique Student Codes", int(profiles["Student Code"].nunique()), fail=len(profiles) != profiles["Student Code"].nunique()),
        _quality_check("Duplicate Student Codes", duplicate_codes, fail=duplicate_codes != 0),
        _quality_check("Spring 2026 leakage rows", int(spring_2026_leakage), fail=int(spring_2026_leakage) != 0),
        _quality_check("Missing Probation Status", int(profiles["Probation Status"].isna().sum()), fail=bool(profiles["Probation Status"].isna().any())),
        _quality_check("Missing Academic Risk Category", int(profiles["Academic Risk Category"].isna().sum()), fail=bool(profiles["Academic Risk Category"].isna().any())),
        _quality_check("Missing latest semester", int(profiles["Latest Semester"].isna().sum()), review=bool(profiles["Latest Semester"].isna().any())),
        _quality_check("Missing Total Credits Earned", int(profiles["Total Credits Earned"].isna().sum()), review=bool(profiles["Total Credits Earned"].isna().any())),
        _quality_check("Negative earned credits", _negative_count(profiles, "Total Credits Earned"), fail=_negative_count(profiles, "Total Credits Earned") != 0),
        _quality_check("Negative attempted credits", _negative_count(profiles, "Total Credits Attempted"), fail=_negative_count(profiles, "Total Credits Attempted") != 0),
        _quality_check("Semester earned credits above semester attempted credits", _scope_credit_count(profiles, "Semester Earned Above Attempted Rows"), fail=_scope_credit_count(profiles, "Semester Earned Above Attempted Rows") != 0),
        _quality_check("Historical earned sum above historical attempted sum", _scope_credit_count(profiles, "Historical Earned Above Attempted Sum"), fail=_scope_credit_count(profiles, "Historical Earned Above Attempted Sum") != 0),
        _quality_check(
            "Earned credits official-vs-derived mismatches",
            int(profiles["Earned Credits Validation Status"].eq(EARNED_MISMATCH).sum()) if "Earned Credits Validation Status" in profiles.columns else 0,
            review="Earned Credits Validation Status" in profiles.columns and bool(profiles["Earned Credits Validation Status"].eq(EARNED_MISMATCH).any()),
        ),
        _quality_check(
            "Attempted credits decrease across semesters",
            int(profiles["Credit Review Note"].fillna("").astype(str).str.contains("decrease across historical semesters").sum()) if "Credit Review Note" in profiles.columns else 0,
            review="Credit Review Note" in profiles.columns and bool(profiles["Credit Review Note"].fillna("").astype(str).str.contains("decrease across historical semesters").any()),
        ),
        _quality_check("Latest SGPA outside 0-4", _outside_gpa_count(profiles, "Latest SGPA"), fail=_outside_gpa_count(profiles, "Latest SGPA") != 0),
        _quality_check("Latest LCGPA outside 0-4", _outside_gpa_count(profiles, "Latest LCGPA"), fail=_outside_gpa_count(profiles, "Latest LCGPA") != 0),
        _quality_check("Latest CGPA outside 0-4", _outside_gpa_count(profiles, "Latest CGPA"), fail=_outside_gpa_count(profiles, "Latest CGPA") != 0),
    ]
    return pd.DataFrame(checks, columns=["Check", "Result", "Status"])


def _compare_earned_credit(official: object, derived: object, calculation_status: object) -> tuple[str, object]:
    if calculation_status != "Calculated" or classify_numeric_value(derived, minimum=None, maximum=None) != "ok":
        return EARNED_INSUFFICIENT, pd.NA
    if classify_numeric_value(official, minimum=None, maximum=None) != "ok":
        return EARNED_NO_OFFICIAL, pd.NA
    difference = float(derived) - float(official)
    if abs(difference) < 0.001:
        return EARNED_MATCHED, 0.0
    return EARNED_MISMATCH, difference


def _quality_check(check: str, result: int, *, fail: bool = False, review: bool = False) -> dict[str, object]:
    if fail:
        status = "FAIL"
    elif review:
        status = "REVIEW"
    else:
        status = "PASS"
    return {"Check": check, "Result": result, "Status": status}


def _scope_credit_count(frame: pd.DataFrame, column: str) -> int:
    if column not in frame.columns:
        return 0
    values = pd.to_numeric(frame[column], errors="coerce").fillna(0)
    return int(values.sum())


def _negative_count(frame: pd.DataFrame, column: str) -> int:
    return sum(classify_numeric_value(value, minimum=0, maximum=None) == "negative" for value in frame[column])


def _outside_gpa_count(frame: pd.DataFrame, column: str) -> int:
    return sum(classify_numeric_value(value, minimum=0, maximum=4) not in {"ok", "blank"} for value in frame[column])


def _historical_semester_frame(semester_summary: pd.DataFrame) -> pd.DataFrame:
    """Return semester rows that are already on the historical side of the holdout."""
    classified = semester_summary if "Partition" in semester_summary.columns else _reject_holdout(semester_summary, "semester summary")
    if not classified["Partition"].eq("Before Spring 2026").all():
        raise RuntimeError("Semester credit input contains a non-historical semester.")
    labels = classified["Semester"].astype(str)
    if labels.eq(HOLDOUT_SEMESTER).any() or labels.str.contains("2026").any():
        raise RuntimeError(f"{HOLDOUT_SEMESTER} entered a historical credit calculation.")
    return classified


def build_student_profile_table(
    course_status: pd.DataFrame,
    snapshots: pd.DataFrame,
    attributes: pd.DataFrame,
) -> pd.DataFrame:
    """Build one profile row per student from course status and the latest snapshot."""
    if course_status.duplicated(["Student Code", "Course Code"]).any():
        raise RuntimeError("Course status has more than one row for a student and course.")
    status_students = set(course_status["Student Code"])
    snapshot_students = set(snapshots["Student Code"])
    if status_students != snapshot_students:
        raise RuntimeError("Course status students and latest snapshots do not match.")
    profiles = snapshots.copy()
    aggregated = _aggregate_course_lists(course_status)
    profiles = profiles.merge(aggregated, on="Student Code", how="left", validate="one_to_one")
    profiles = profiles.merge(attributes, on="Student Code", how="left", validate="one_to_one")
    missing_source = profiles["Specialization Source"].isna()
    profiles.loc[missing_source, "Current Specialization"] = "Not available before Spring 2026"
    profiles.loc[missing_source, "Specialization Source"] = "Not available before Spring 2026"
    profiles.loc[missing_source, "Attribute Review Reason"] = "No specialization field exists on a pre-Spring 2026 snapshot"
    unresolved_specialization = profiles["Current Specialization"].isna()
    profiles.loc[unresolved_specialization, "Current Specialization"] = "Manual Review"
    profiles["Current Level"] = profiles["Current Level Raw"].map(_level_label)
    profiles["Total Credits Attempted"] = profiles["Official Credits Attempted"]
    profiles["Total Credits Earned"] = profiles["Official Credits Earned"]
    profiles["Attempted Credits Difference"] = [
        _numeric_difference(official, reconstructed)
        for official, reconstructed in zip(profiles["Official Credits Attempted"], profiles["Reconstructed Credits Attempted"])
    ]
    profiles["Earned Credits Difference"] = [
        _numeric_difference(official, reconstructed)
        for official, reconstructed in zip(profiles["Official Credits Earned"], profiles["Reconstructed Credits Earned"])
    ]
    profiles["Remaining Courses Count"] = "Pending"
    profiles["Remaining Courses Status"] = "Pending Study Plan Matching"
    probation = profiles.apply(_probation_from_row, axis=1, result_type="expand")
    probation.columns = ["Probation Status", "Is Probation", "Probation Review Reason"]
    profiles = pd.concat([profiles, probation], axis=1)
    risk = profiles["Probation Status"].map(_risk_category)
    profiles["Academic Risk Category"] = risk
    profiles["Academic Risk Basis"] = ACADEMIC_RISK_BASIS
    profiles["Risk Rule Version"] = RISK_RULE_VERSION
    profiles["Credits Validation Status"] = CREDITS_NOTE
    profiles["Cutoff Semester"] = CUTOFF_SEMESTER
    profiles["Task 7 Logic Version"] = TASK7_LOGIC_VERSION
    review = profiles.apply(_profile_review, axis=1, result_type="expand")
    review.columns = ["Profile Validation Status", "Manual Review Required", "Manual Review Reasons"]
    profiles = pd.concat([profiles, review], axis=1)
    profiles = profiles.sort_values("Student Code", kind="mergesort").reset_index(drop=True)
    if not profiles["Student Code"].is_unique:
        raise RuntimeError('profile["Student Code"].is_unique == True failed.')
    return _order_columns(profiles, PROFILE_COLUMNS)


def build_supporting_tables(course_status: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Split course status into completed, failed, withdrawn, and repeated rows."""
    completed_columns = [
        "Student Code", "Course Code", "Course Name", "Completion Semester", "Completion Grade",
        "Completion Grade Point", "Minimum Passing Grade", "Minimum Passing Grade Point",
        "Number of Attempts Before Completion", "Passed On First Attempt", "Eventually Passed",
    ]
    unresolved_columns = [
        "Student Code", "Course Code", "Course Name", "Attempt Count", "Latest Attempt Semester",
        "Latest Grade", "Latest Grade Point", "Minimum Passing Grade", "Ever Failed",
        "Ever Withdrawn", "Eventually Passed", "Final Status",
    ]
    repeated_columns = [
        "Student Code", "Course Code", "Course Name", "Attempt Count", "First Attempt Semester",
        "Latest Attempt Semester", "First Grade", "Latest Grade", "Ever Failed", "Ever Withdrawn",
        "Eventually Passed", "Final Status",
    ]
    return {
        "completed": _subset(course_status, "Is Completed", completed_columns),
        "failed": _subset(course_status, "Is Current Failed", unresolved_columns).rename(columns={"Attempt Count": "Number of Attempts", "Minimum Passing Grade": "Required Passing Grade"}),
        "withdrawn": _subset(course_status, "Is Current Withdrawn", unresolved_columns).rename(columns={"Attempt Count": "Number of Attempts", "Minimum Passing Grade": "Required Passing Grade"}),
        "repeated": _subset(course_status, "Is Repeated", repeated_columns),
    }


def build_profile_validation_table(
    profiles: pd.DataFrame,
    course_status: pd.DataFrame,
) -> pd.DataFrame:
    """Collect student and course issues that a reviewer needs to see."""
    rows: list[dict[str, object]] = []
    for profile in profiles.to_dict(orient="records"):
        reasons = profile["Manual Review Reasons"]
        if _missing(reasons):
            continue
        for reason in [part.strip() for part in str(reasons).split("|") if part.strip()]:
            rows.append({
                "Student Code": profile["Student Code"],
                "Course Code": pd.NA,
                "Severity": "Manual Review",
                "Issue": reason,
            })
    for status in course_status.to_dict(orient="records"):
        notes = status.get("Review Notes")
        if _missing(notes):
            continue
        for note in [part.strip() for part in str(notes).split("|") if part.strip()]:
            severity = "Warning" if note.startswith("Warning:") else "Manual Review"
            rows.append({
                "Student Code": status["Student Code"],
                "Course Code": status["Course Code"],
                "Severity": severity,
                "Issue": note.removeprefix("Warning: ").strip(),
            })
    validation = pd.DataFrame(rows, columns=["Student Code", "Course Code", "Severity", "Issue"])
    if validation.empty:
        return validation
    return validation.sort_values(["Student Code", "Course Code", "Issue"], kind="mergesort").reset_index(drop=True)


def validate_student_profiles(
    profiles: pd.DataFrame,
    course_status: pd.DataFrame,
    *,
    spring_2026_leakage: int,
) -> dict[str, object]:
    """Raise when a profile integrity rule fails. Return the passing counts."""
    if spring_2026_leakage != 0:
        raise RuntimeError(f"Spring 2026 leakage is {spring_2026_leakage}, expected 0.")
    if not profiles["Student Code"].is_unique:
        raise RuntimeError('profile["Student Code"].is_unique == True failed.')
    if len(profiles) != profiles["Student Code"].nunique():
        raise RuntimeError("Profile row count does not equal the number of student codes.")
    if profiles["Latest Semester"].astype(str).str.contains("2026").any():
        raise RuntimeError("A profile latest semester contains 2026.")
    if course_status["Latest Attempt Semester"].astype(str).str.contains("2026").any():
        raise RuntimeError("A course status semester contains 2026.")
    completed_and_failed = course_status["Is Completed"] & course_status["Is Current Failed"]
    completed_and_withdrawn = course_status["Is Completed"] & course_status["Is Current Withdrawn"]
    if completed_and_failed.any() or completed_and_withdrawn.any():
        raise RuntimeError("A course is both completed and currently failed or withdrawn.")
    repeated_mismatch = course_status["Is Repeated"] != course_status["Attempt Count"].gt(1)
    if repeated_mismatch.any():
        raise RuntimeError("Repeated status does not match attempt count.")
    if not course_status.loc[course_status["Eventually Passed"], "Is Completed"].all():
        raise RuntimeError("A later or earlier pass was not classified as completed.")
    if course_status.duplicated(["Student Code", "Course Code"]).any():
        raise RuntimeError("Course status is not unique by student and course.")
    for flag, count_column in (
        ("Is Completed", "Completed Course Count"),
        ("Is Current Failed", "Failed Course Count"),
        ("Is Current Withdrawn", "Withdrawn Course Count"),
        ("Is Repeated", "Repeated Course Count"),
    ):
        if int(course_status[flag].sum()) != int(profiles[count_column].sum()):
            raise RuntimeError(f"{count_column} does not match the course-status table.")
    _assert_list_counts(profiles, "Completed Course Count", "Completed Courses List")
    _assert_list_counts(profiles, "Failed Course Count", "Failed Courses List")
    _assert_list_counts(profiles, "Withdrawn Course Count", "Withdrawn Courses List")
    _assert_list_counts(profiles, "Repeated Course Count", "Repeated Courses List")
    _assert_non_negative(profiles, "Total Credits Attempted")
    _assert_non_negative(profiles, "Total Credits Earned")
    _assert_non_negative(profiles, "Reconstructed Credits Attempted")
    _assert_non_negative(profiles, "Reconstructed Credits Earned")
    for column in ("Latest SGPA", "Latest LCGPA", "Latest CGPA"):
        invalid = [
            classify_numeric_value(value, minimum=0, maximum=4)
            for value in profiles[column]
            if classify_numeric_value(value, minimum=0, maximum=4) not in {"ok", "blank"}
        ]
        if invalid:
            raise RuntimeError(f"{column} contains a value outside the existing 0 to 4 check: {invalid[0]}.")
    return {"passed": True, "spring_2026_leakage": 0}


def build_summary_sheet(
    profiles: pd.DataFrame,
    course_status: pd.DataFrame,
    validation: pd.DataFrame,
    boundary: Mapping[str, object],
    *,
    run_timestamp: str,
) -> pd.DataFrame:
    """Build the workbook summary, including the required Task 7 counts."""
    probation = profiles["Probation Status"].value_counts()
    presented_status = _presentation_notes(profiles, course_status)["Profile Status"].value_counts()
    missing_rule_mask = course_status["Review Notes"].astype(str).str.contains("No passing-grade rule", na=False)
    missing_rule_courses = course_status.loc[missing_rule_mask, "Course Code"]
    manual_students = set(profiles.loc[profiles["Manual Review Required"].eq("Yes"), "Student Code"])
    metrics = [
        ("Number of students", int(len(profiles))),
        ("Number of profiles created", int(len(profiles))),
        ("Number of unique Student Codes", int(profiles["Student Code"].nunique())),
        ("Number of completed student-course pairs", int(course_status["Is Completed"].sum())),
        ("Number of unresolved failed student-course pairs", int(course_status["Is Current Failed"].sum())),
        ("Number of unresolved withdrawn student-course pairs", int(course_status["Is Current Withdrawn"].sum())),
        ("Number of repeated student-course pairs", int(course_status["Is Repeated"].sum())),
        ("Number of probation students", int(probation.get("Probation", 0))),
        ("Number of non-probation students", int(probation.get("Not Probation", 0))),
        ("Number requiring probation manual review", int(probation.get("Manual Review", 0))),
        ("Number requiring any profile manual review", int(len(manual_students))),
        ("Profile validation rows", int(len(validation))),
        ("Number of courses missing passing-grade rules", int(missing_rule_courses.nunique())),
        ("Student-course pairs missing a passing-grade rule", int(missing_rule_mask.sum())),
        ("Number of profile rows with missing latest SGPA", int(profiles["Latest SGPA"].isna().sum())),
        ("Number with missing latest LCGPA", int(profiles["Latest LCGPA"].isna().sum())),
        ("Number with missing latest CGPA", int(profiles["Latest CGPA"].isna().sum())),
        ("Number with missing level", int(profiles["Current Level"].eq("Unknown/Manual Review").sum())),
        ("Number with missing specialization", int(profiles["Current Specialization"].eq("Not available before Spring 2026").sum())),
        ("Profiles with status Valid", int(presented_status.get("Valid", 0))),
        ("Profiles with status Valid with Missing Information", int(presented_status.get("Valid with Missing Information", 0))),
        ("Profiles with status Manual Review Required", int(presented_status.get("Manual Review Required", 0))),
        ("Historical cutoff semester", CUTOFF_SEMESTER),
        ("Spring 2026 leakage count", int(boundary["spring_2026_leakage"])),
        ("Historical course rows used", int(boundary["historical_course_rows"])),
        ("Historical semester rows used", int(boundary["historical_semester_rows"])),
        ("Historical students", int(boundary["historical_students"])),
        ("Earliest semester", boundary["earliest_semester"]),
        ("Latest historical semester", boundary["latest_semester"]),
        ("Spring 2026 course rows excluded", int(boundary["spring_2026_course_rows_excluded"])),
        ("Spring 2026 semester rows excluded", int(boundary["spring_2026_semester_rows_excluded"])),
        ("Spring 2026 students excluded from feature construction", int(boundary["spring_2026_students_excluded_from_features"])),
        ("Spring 2026 students with no historical profile", int(boundary["spring_2026_students_without_historical_profile"])),
        ("Unresolved transcript-failure pairs", int(course_status["Final Status"].eq("Failed").sum())),
        ("Unresolved below-minimum pairs", int(course_status["Final Status"].eq("Below Passing Grade").sum())),
        ("Task 7 logic version", TASK7_LOGIC_VERSION),
        ("Risk rule version", RISK_RULE_VERSION),
        ("Run timestamp", run_timestamp),
        ("Grade-rule source", "outputs/course_scope_preparation.xlsx / Passing Grades"),
        ("Course-history source", "outputs/course_scope_preparation.xlsx / Validated History"),
        ("Semester-summary source", "outputs/historical_data_before_spring_2026.xlsx / Semester Summary"),
        ("Credits assumption", CREDITS_NOTE),
        ("Total Credits Attempted Source", ATTEMPTED_CREDITS_SOURCE),
        ("Total Credits Earned Source", EARNED_CREDITS_SOURCE),
        ("Students with reconstructed Total Credits Earned", int(profiles["Total Credits Earned"].notna().sum())),
        ("Students missing Total Credits Earned", int(profiles["Total Credits Earned"].isna().sum())),
        ("Students with official historical earned-credit comparison", int(profiles["Official Credits Earned"].notna().sum()) if "Official Credits Earned" in profiles.columns else 0),
        ("Matched official historical earned-credit comparisons", int(profiles["Earned Credits Validation Status"].eq(EARNED_MATCHED).sum()) if "Earned Credits Validation Status" in profiles.columns else 0),
        ("Mismatched earned-credit comparisons", int(profiles["Earned Credits Validation Status"].eq(EARNED_MISMATCH).sum()) if "Earned Credits Validation Status" in profiles.columns else 0),
        ("Risk assumption", ACADEMIC_RISK_BASIS),
        ("Remaining-course assumption", "Remaining Courses Count compares completed historical courses with the official study plan for the student's level and pathway. A passed course is removed. A failed, withdrawn, or not-yet-taken required course stays. An elective slot is filled by any completed course from that slot's elective pool."),
    ]
    return pd.DataFrame(metrics, columns=["Metric", "Value"])


def format_task7_report(summary: pd.DataFrame, *, passed: bool) -> str:
    """Return the compact end-of-task report."""
    values = dict(zip(summary["Metric"], summary["Value"]))
    result = "PASS" if passed else "FAIL"
    return "\n".join([
        "TASK 7 — STUDENT ACADEMIC PROFILE CREATION",
        "------------------------------------------",
        f"Historical cutoff: {values['Historical cutoff semester']}",
        f"Spring 2026 leakage: {values['Spring 2026 leakage count']}",
        f"Students processed: {values['Historical students']}",
        f"Profiles created: {values['Number of profiles created']}",
        f"Unique profile students: {values['Number of unique Student Codes']}",
        "",
        f"Completed student-course pairs: {values['Number of completed student-course pairs']}",
        f"Current unresolved failures: {values['Number of unresolved failed student-course pairs']}",
        f"Current unresolved withdrawals: {values['Number of unresolved withdrawn student-course pairs']}",
        f"Repeated student-course pairs: {values['Number of repeated student-course pairs']}",
        "",
        f"Probation: {values['Number of probation students']}",
        f"Not Probation: {values['Number of non-probation students']}",
        f"Probation Manual Review: {values['Number requiring probation manual review']}",
        "",
        f"Profile status Valid: {values['Profiles with status Valid']}",
        f"Profile status Valid with Missing Information: {values['Profiles with status Valid with Missing Information']}",
        f"Profile status Manual Review Required: {values['Profiles with status Manual Review Required']}",
        f"Students with missing pre-holdout specialization: {values['Number with missing specialization']}",
        f"Courses missing passing-grade rule: {values['Number of courses missing passing-grade rules']}",
        "",
        f"Total Credits Earned populated: {values['Students with reconstructed Total Credits Earned']}",
        f"Total Credits Earned missing: {values['Students missing Total Credits Earned']}",
        f"Official historical earned-credit comparisons: {values['Students with official historical earned-credit comparison']}",
        f"Matched official comparisons: {values['Matched official historical earned-credit comparisons']}",
        f"Mismatched official comparisons: {values['Mismatched earned-credit comparisons']}",
        "",
        "Output:",
        "student_academic_profiles.xlsx",
        "",
        f"Task 7 validation: {result}",
    ])


COMMON_YEAR_PATH = "Common Year 1"
SOFTWARE_YEAR2_PATH = "Software Engineering Year 2"
NETWORK_YEAR2_PATH = "Network Computing and Security Year 2"
GENERAL_ELECTIVE_POOL = "General Requirement"
PATHWAY_YEAR2 = {
    "Software Engineering": SOFTWARE_YEAR2_PATH,
    "Data Science and Artificial Intelligence": SOFTWARE_YEAR2_PATH,
    "Cyber and Information Security": NETWORK_YEAR2_PATH,
}
PLAN_LEVELS = ("Diploma", "Advanced Diploma", "Bachelor")


def assign_remaining_course_counts(
    profiles: pd.DataFrame,
    course_status: pd.DataFrame,
    study_plan: pd.DataFrame,
    elective_pools: pd.DataFrame,
) -> pd.DataFrame:
    """Fill Remaining Courses Count from the official study plan.

    A completed course is removed. Failed, withdrawn, and not-yet-taken
    required courses remain. One completed course from the correct elective
    pool fills one elective slot. The student profile's other fields are not
    changed. Spring 2026 rows are not read here; course status is already
    limited to the historical cutoff.
    """
    catalog = _study_plan_catalog(study_plan, elective_pools)
    completed = _codes_by_student(course_status, completed_only=True)
    attempted = _codes_by_student(course_status, completed_only=False)
    counts: list[object] = []
    explanations: list[str] = []
    for profile in profiles.to_dict(orient="records"):
        student_code = profile["Student Code"]
        count, explanation = _remaining_for_student(
            level=_canonical_level(profile.get("Current Level")),
            specialization=_known_specialization(profile.get("Current Specialization")),
            completed=completed.get(student_code, set()),
            attempted=attempted.get(student_code, set()),
            catalog=catalog,
        )
        counts.append(count)
        explanations.append(explanation)
    updated = profiles.copy()
    updated["Remaining Courses Count"] = counts
    updated["Remaining Courses Status"] = explanations
    numeric_counts = pd.to_numeric(updated["Remaining Courses Count"], errors="coerce")
    if numeric_counts.dropna().lt(0).any():
        raise RuntimeError("A remaining-course count is negative.")
    return updated


def _study_plan_catalog(study_plan: pd.DataFrame, elective_pools: pd.DataFrame) -> dict[str, object]:
    plan = study_plan.copy()
    plan["Course Code"] = plan["Course Code"].map(normalize_course_code)
    plan["Core"] = plan["Is Core"].map(_yes_text)
    plan["Elective"] = plan["Is Elective"].map(_yes_text)
    cores: dict[tuple[str, str], set[str]] = {}
    slots: dict[tuple[str, str], list[str]] = {}
    for (path, level), group in plan.groupby(["Specialization Path", "Level"], sort=False):
        key = (str(path), str(level))
        cores[key] = set(group.loc[group["Core"] & group["Course Code"].notna(), "Course Code"].astype(str))
        elective_rows = group.loc[group["Elective"]]
        slots[key] = [str(pool) for pool in elective_rows["Elective Pool"].dropna()]
    pools: dict[str, set[str]] = {}
    pool_frame = elective_pools.copy()
    pool_frame["Course Code"] = pool_frame["Course Code"].map(normalize_course_code)
    for pool_name, group in pool_frame.groupby("Elective Pool", sort=False):
        pools[str(pool_name)] = set(group["Course Code"].dropna().astype(str))
    all_cores = set().union(*cores.values()) if cores else set()
    return {"cores": cores, "slots": slots, "pools": pools, "all_cores": all_cores}


def _remaining_for_student(
    *,
    level: str | None,
    specialization: str | None,
    completed: set[str],
    attempted: set[str],
    catalog: dict[str, object],
) -> tuple[object, str]:
    if level is None:
        level = _level_from_attempted_courses(attempted, catalog)
        level_note = "level inferred from attempted plan courses"
    else:
        level_note = "official level"
    if level not in PLAN_LEVELS:
        return pd.NA, "Remaining courses were not calculated because the level is not on the study plan"
    candidates = _pathway_candidates(level, specialization, attempted, catalog)
    if not candidates:
        return pd.NA, "Remaining courses were not calculated because the study-plan pathway is not resolved"
    counted = [
        (
            _remaining_count(level, year2, spec, completed, catalog),
            year2,
            spec,
        )
        for year2, spec in candidates
    ]
    unique_counts = {item[0] for item in counted}
    if len(unique_counts) == 1:
        count = next(iter(unique_counts))
        pathway = _pathway_label(level, counted[0][1], counted[0][2] if len(candidates) == 1 else None)
        if len(candidates) > 1:
            pathway = f"{level}; every allowed pathway leaves this same count"
        return int(count), f"Calculated from the study plan ({level_note}): {pathway}"
    best = max(counted, key=lambda item: (_pathway_evidence(item[1], item[2], attempted, catalog), -item[0]))
    evidence = [
        _pathway_evidence(year2, spec, attempted, catalog)
        for _, year2, spec in counted
    ]
    if evidence.count(max(evidence)) > 1:
        return pd.NA, "Remaining courses were not calculated because more than one study-plan pathway fits the courses"
    count, year2, spec = best
    return int(count), f"Calculated from the study plan ({level_note}): {_pathway_label(level, year2, spec)}"


def _pathway_candidates(
    level: str,
    specialization: str | None,
    attempted: set[str],
    catalog: dict[str, object],
) -> list[tuple[str | None, str | None]]:
    if specialization in PATHWAY_YEAR2:
        return [(PATHWAY_YEAR2[specialization], None if level == "Diploma" else specialization)]
    year2 = _year2_from_attempts(attempted, catalog)
    if level == "Diploma":
        return [(year2, None)]
    evidenced = _specialization_from_attempts(attempted, catalog)
    if evidenced:
        return [(PATHWAY_YEAR2[evidenced], evidenced)]
    if year2 == NETWORK_YEAR2_PATH:
        return [(year2, "Cyber and Information Security")]
    if year2 == SOFTWARE_YEAR2_PATH:
        options = ["Software Engineering", "Data Science and Artificial Intelligence"]
        chosen = _elective_specialization(attempted, catalog, options)
        if chosen:
            return [(year2, chosen)]
        return [(year2, name) for name in options]
    return [
        (SOFTWARE_YEAR2_PATH, "Software Engineering"),
        (SOFTWARE_YEAR2_PATH, "Data Science and Artificial Intelligence"),
        (NETWORK_YEAR2_PATH, "Cyber and Information Security"),
    ]


def _remaining_count(
    level: str,
    year2: str | None,
    specialization: str | None,
    completed: set[str],
    catalog: dict[str, object],
) -> int:
    cores, slots = _required_parts(level, year2, specialization, catalog)
    if year2 is None and level == "Diploma":
        shared = _core_set(COMMON_YEAR_PATH, "Diploma", catalog) | (
            _core_set(SOFTWARE_YEAR2_PATH, "Diploma", catalog) & _core_set(NETWORK_YEAR2_PATH, "Diploma", catalog)
        )
        remaining = len(shared - completed) + 5
        slots = [GENERAL_ELECTIVE_POOL]
        cores = shared
    else:
        remaining = len(cores - completed)
    used: set[str] = set()
    pools = catalog["pools"]
    for pool_name in slots:
        if pool_name == GENERAL_ELECTIVE_POOL:
            options = _general_elective_options(completed, used, catalog)
        else:
            options = (completed & pools.get(pool_name, set())) - cores - used
        if options:
            used.add(sorted(options)[0])
        else:
            remaining += 1
    return int(remaining)


def _required_parts(
    level: str,
    year2: str | None,
    specialization: str | None,
    catalog: dict[str, object],
) -> tuple[set[str], list[str]]:
    cores = set(_core_set(COMMON_YEAR_PATH, "Diploma", catalog))
    slots: list[str] = []
    if year2:
        cores |= _core_set(year2, "Diploma", catalog)
        slots.extend(_slot_list(year2, "Diploma", catalog))
    if level in {"Advanced Diploma", "Bachelor"} and specialization:
        cores |= _core_set(specialization, "Advanced Diploma", catalog)
        slots.extend(_slot_list(specialization, "Advanced Diploma", catalog))
    if level == "Bachelor" and specialization:
        cores |= _core_set(specialization, "Bachelor", catalog)
        slots.extend(_slot_list(specialization, "Bachelor", catalog))
    return cores, slots


def _general_elective_options(completed: set[str], used: set[str], catalog: dict[str, object]) -> set[str]:
    outside_plan = completed - catalog["all_cores"] - used
    pool_codes = set().union(*catalog["pools"].values()) if catalog["pools"] else set()
    preferred = outside_plan - pool_codes
    return preferred or outside_plan


def _year2_from_attempts(attempted: set[str], catalog: dict[str, object]) -> str | None:
    software = _core_set(SOFTWARE_YEAR2_PATH, "Diploma", catalog)
    network = _core_set(NETWORK_YEAR2_PATH, "Diploma", catalog)
    software_hits = attempted & (software - network)
    network_hits = attempted & (network - software)
    if software_hits and not network_hits:
        return SOFTWARE_YEAR2_PATH
    if network_hits and not software_hits:
        return NETWORK_YEAR2_PATH
    if len(software_hits) > len(network_hits):
        return SOFTWARE_YEAR2_PATH
    if len(network_hits) > len(software_hits):
        return NETWORK_YEAR2_PATH
    return None


def _specialization_from_attempts(attempted: set[str], catalog: dict[str, object]) -> str | None:
    scores = {
        specialization: len(attempted & _specialization_signature(specialization, catalog))
        for specialization in PATHWAY_YEAR2
    }
    best_score = max(scores.values(), default=0)
    if best_score == 0:
        return None
    winners = [name for name, score in scores.items() if score == best_score]
    if len(winners) == 1:
        return winners[0]
    return None


def _specialization_signature(specialization: str, catalog: dict[str, object]) -> set[str]:
    advanced = _core_set(specialization, "Advanced Diploma", catalog)
    bachelor = _core_set(specialization, "Bachelor", catalog)
    others = set()
    for other in PATHWAY_YEAR2:
        if other == specialization:
            continue
        others |= _core_set(other, "Advanced Diploma", catalog)
        others |= _core_set(other, "Bachelor", catalog)
    return (advanced | bachelor) - others


def _pathway_evidence(year2: str | None, specialization: str | None, attempted: set[str], catalog: dict[str, object]) -> int:
    evidence = 0
    if year2:
        own = _core_set(year2, "Diploma", catalog)
        other = _core_set(
            NETWORK_YEAR2_PATH if year2 == SOFTWARE_YEAR2_PATH else SOFTWARE_YEAR2_PATH,
            "Diploma",
            catalog,
        )
        evidence += len(attempted & (own - other)) * 3
    if specialization:
        evidence += len(attempted & _specialization_signature(specialization, catalog)) * 10
        evidence += _exclusive_elective_count(attempted, catalog, specialization, list(PATHWAY_YEAR2))
    return evidence


def _elective_specialization(attempted: set[str], catalog: dict[str, object], options: list[str]) -> str | None:
    scores = {
        name: _exclusive_elective_count(attempted, catalog, name, options)
        for name in options
    }
    best_score = max(scores.values(), default=0)
    if best_score == 0:
        return None
    winners = [name for name, score in scores.items() if score == best_score]
    if len(winners) == 1:
        return winners[0]
    return None


def _exclusive_elective_count(
    attempted: set[str],
    catalog: dict[str, object],
    specialization: str,
    options: list[str],
) -> int:
    pools = catalog["pools"]
    own = set(pools.get(_major_pool(specialization, catalog) or "", set()))
    others: set[str] = set()
    for name in options:
        if name == specialization:
            continue
        others |= set(pools.get(_major_pool(name, catalog) or "", set()))
    return len(attempted & (own - others))


def _major_pool(specialization: str, catalog: dict[str, object]) -> str | None:
    for level in ("Advanced Diploma", "Bachelor"):
        for slot in _slot_list(specialization, level, catalog):
            if slot != GENERAL_ELECTIVE_POOL:
                return slot
    return None


def _level_from_attempted_courses(attempted: set[str], catalog: dict[str, object]) -> str | None:
    higher = set()
    advanced = set()
    for specialization in PATHWAY_YEAR2:
        higher |= _core_set(specialization, "Bachelor", catalog)
        advanced |= _core_set(specialization, "Advanced Diploma", catalog)
    if attempted & higher:
        return "Bachelor"
    if attempted & advanced:
        return "Advanced Diploma"
    if attempted:
        return "Diploma"
    return None


def _core_set(path: str, level: str, catalog: dict[str, object]) -> set[str]:
    return set(catalog["cores"].get((path, level), set()))


def _slot_list(path: str, level: str, catalog: dict[str, object]) -> list[str]:
    return list(catalog["slots"].get((path, level), []))


def _codes_by_student(course_status: pd.DataFrame, *, completed_only: bool) -> dict[object, set[str]]:
    frame = course_status.loc[course_status["Is Completed"]] if completed_only else course_status
    grouped: dict[object, set[str]] = {}
    for student_code, group in frame.groupby("Student Code", sort=False):
        grouped[student_code] = set(group["Course Code"].dropna().astype(str))
    return grouped


def _canonical_level(value: object) -> str | None:
    if _missing(value):
        return None
    text = str(value).strip().casefold()
    for level in PLAN_LEVELS:
        if text == level.casefold():
            return level
    return None


def _known_specialization(value: object) -> str | None:
    if _missing(value):
        return None
    text = str(value).strip()
    if text in PATHWAY_YEAR2:
        return text
    return None


def _pathway_label(level: str, year2: str | None, specialization: str | None) -> str:
    parts = [level]
    if year2:
        parts.append(year2)
    if specialization:
        parts.append(specialization)
    return ", ".join(parts)


def _yes_text(value: object) -> bool:
    return str(value).strip().casefold() == "yes"


EXPORTED_PROFILE_COLUMNS = [
    "Student Code",
    "Current Level",
    "Current Specialization",
    "Latest Semester",
    "Latest SGPA",
    "Latest LCGPA",
    "Latest CGPA",
    "Total Credits Attempted",
    "Total Credits Earned",
    "Completed Courses",
    "Failed Courses",
    "Withdrawn Courses",
    "Repeated Courses",
    "Remaining Courses Count",
    "Probation Status",
    "Academic Risk Category",
]

EXPORTED_COURSE_COLUMNS = [
    "Student Code",
    "Course Code",
    "Course Name",
    "Attempt Count",
    "First Attempt Semester",
    "Latest Attempt Semester",
    "First Grade",
    "Latest Grade",
    "Minimum Passing Grade",
    "Ever Failed",
    "Ever Withdrawn",
    "Eventually Passed",
    "Repeated",
    "Final Status",
    "Completion Semester",
    "Completion Grade",
    "Review Note",
]

VALIDATION_METRICS = [
    "Historical cutoff semester",
    "Historical students",
    "Number of profiles created",
    "Number of unique Student Codes",
    "Spring 2026 leakage count",
    "Historical course rows used",
    "Historical semester rows used",
    "Number of completed student-course pairs",
    "Number of unresolved failed student-course pairs",
    "Number of unresolved withdrawn student-course pairs",
    "Number of repeated student-course pairs",
    "Number of probation students",
    "Number of non-probation students",
    "Number with missing specialization",
    "Number with missing level",
    "Number of courses missing passing-grade rules",
    "Student-course pairs missing a passing-grade rule",
]

VALIDATION_METRIC_LABELS = {
    "Number of profiles created": "Profiles created",
    "Number of unique Student Codes": "Unique Student Codes",
    "Spring 2026 leakage count": "Spring 2026 leakage",
    "Number of completed student-course pairs": "Completed student-course pairs",
    "Number of unresolved failed student-course pairs": "Unresolved failed student-course pairs",
    "Number of unresolved withdrawn student-course pairs": "Unresolved withdrawn student-course pairs",
    "Number of repeated student-course pairs": "Repeated student-course pairs",
    "Number of probation students": "Probation students",
    "Number of non-probation students": "Non-probation students",
    "Number with missing specialization": "Students with missing pre-holdout specialization",
    "Number with missing level": "Students with missing level",
    "Number of courses missing passing-grade rules": "Courses missing passing-grade rules",
    "Student-course pairs missing a passing-grade rule": "Student-course pairs affected by missing passing-grade rules",
}


def present_student_profiles(profiles: pd.DataFrame, course_status: pd.DataFrame) -> pd.DataFrame:
    """Return one student row with only the required profile fields.

    ``course_status`` is accepted so existing callers stay stable. Course
    evidence is exported on the course-history sheet, not on this profile.
    """
    del course_status
    work = profiles.reset_index(drop=True)
    presented = pd.DataFrame({
        "Student Code": work["Student Code"].astype(str),
        "Current Level": work["Current Level"],
        "Current Specialization": work["Current Specialization"],
        "Latest Semester": work["Latest Semester"],
        "Latest SGPA": work["Latest SGPA"],
        "Latest LCGPA": work["Latest LCGPA"],
        "Latest CGPA": work["Latest CGPA"],
        "Total Credits Attempted": work["Total Credits Attempted"],
        "Total Credits Earned": work["Total Credits Earned"],
        "Completed Courses": work["Completed Courses List"],
        "Failed Courses": work["Failed Courses List"],
        "Withdrawn Courses": work["Withdrawn Courses List"],
        "Repeated Courses": work["Repeated Courses List"],
        "Remaining Courses Count": work["Remaining Courses Count"],
        "Probation Status": work["Probation Status"],
        "Academic Risk Category": work["Academic Risk Category"],
    })
    return presented.loc[:, EXPORTED_PROFILE_COLUMNS]


def present_course_history(course_status: pd.DataFrame) -> pd.DataFrame:
    """Return one readable row per student and course."""
    presented = pd.DataFrame({
        "Student Code": course_status["Student Code"].astype(str),
        "Course Code": course_status["Course Code"].astype(str),
        "Course Name": course_status["Course Name"],
        "Attempt Count": course_status["Attempt Count"],
        "First Attempt Semester": course_status["First Attempt Semester"],
        "Latest Attempt Semester": course_status["Latest Attempt Semester"],
        "First Grade": course_status["First Grade"],
        "Latest Grade": course_status["Latest Grade"],
        "Minimum Passing Grade": course_status["Minimum Passing Grade"],
        "Ever Failed": _yes_no(course_status["Ever Failed"]),
        "Ever Withdrawn": _yes_no(course_status["Ever Withdrawn"]),
        "Eventually Passed": _yes_no(course_status["Eventually Passed"]),
        "Repeated": _yes_no(course_status["Is Repeated"]),
        "Final Status": course_status["Final Status"],
        "Completion Semester": course_status["Completion Semester"],
        "Completion Grade": course_status["Completion Grade"],
        "Review Note": course_status["Review Notes"].map(_short_course_note),
    })
    return presented.loc[:, EXPORTED_COURSE_COLUMNS].reset_index(drop=True)


def present_validation_tables(
    profiles: pd.DataFrame,
    course_status: pd.DataFrame,
    summary: pd.DataFrame,
    *,
    validation_status: str,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return the concise metric table and the meaningful issue rows."""
    values = dict(zip(summary["Metric"], summary["Value"]))
    metric_rows = [
        (VALIDATION_METRIC_LABELS.get(metric, metric), values[metric])
        for metric in VALIDATION_METRICS
    ]
    metric_rows.append(("Task 7 validation status", validation_status))
    metrics = pd.DataFrame(metric_rows, columns=["Metric", "Value"])
    issues = _presentation_issues(profiles, course_status)
    return metrics, issues


def export_student_profile_outputs(
    profiles: pd.DataFrame,
    course_status: pd.DataFrame,
    supporting: Mapping[str, pd.DataFrame],
    validation: pd.DataFrame,
    summary: pd.DataFrame,
    workbook_path: Path,
    *,
    validation_status: str = "PASS",
) -> None:
    """Write the student-profile workbook.

    ``supporting`` and ``validation`` remain available to callers. The workbook
    presents their evidence in Course History Status and Validation Summary.
    """
    del supporting, validation
    presented_profiles = present_student_profiles(profiles, course_status)
    presented_courses = present_course_history(course_status)
    metrics, issues = present_validation_tables(
        profiles,
        course_status,
        summary,
        validation_status=validation_status,
    )
    audit = build_earned_credit_audit(profiles)
    leakage = dict(zip(summary["Metric"], summary["Value"])).get("Spring 2026 leakage count", 0)
    checks = build_profile_quality_checks(profiles, spring_2026_leakage=int(leakage))

    def write_workbook(path: Path) -> None:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            presented_profiles.to_excel(writer, sheet_name="Student Academic Profiles", index=False)
            presented_courses.to_excel(writer, sheet_name="Course History Status", index=False)
            metrics.to_excel(writer, sheet_name="Validation Summary", index=False)
            _write_validation_issues(writer.book["Validation Summary"], issues)
            _write_quality_checks(writer.book["Validation Summary"], checks)
            if not audit.empty:
                audit.to_excel(writer, sheet_name="Earned Credit Validation", index=False)
                _format_credit_audit_sheet(writer.book["Earned Credit Validation"])
            _format_profile_sheet(writer.book["Student Academic Profiles"])
            _format_course_sheet(writer.book["Course History Status"])
            _format_validation_sheet(writer.book["Validation Summary"], len(metrics), len(issues))

    write_atomic(workbook_path, write_workbook)


def _presentation_notes(profiles: pd.DataFrame, course_status: pd.DataFrame) -> pd.DataFrame:
    course_notes = _course_note_flags(course_status)
    rows: list[dict[str, str]] = []
    for profile in profiles.to_dict(orient="records"):
        student_code = profile["Student Code"]
        flags = course_notes.get(student_code, {})
        reasons = set(_reason_parts(profile.get("Manual Review Reasons")))
        information: list[str] = []
        review: list[str] = []
        blocking: list[str] = []
        if profile.get("Current Specialization") == "Not available before Spring 2026":
            information.append("Specialization not available before Spring 2026")
        if profile.get("Current Level") == "Unknown/Manual Review":
            information.append("Current level missing on latest historical semester")
        if "Specialization text is not in the known label list" in reasons:
            information.append("Specialization text needs review")
        if flags.get("missing_rule"):
            review.append("One or more courses have no passing-grade rule")
        if flags.get("unclassified"):
            review.append("One or more course outcomes could not be classified")
        if profile.get("Current Specialization") == "Manual Review" or "Transcript specialization conflicts with the inferred specialization path" in reasons:
            blocking.append("Conflicting specialization values")
        if "Latest semester has conflicting summary rows" in reasons or "Probation cannot be decided from conflicting semester rows" in reasons:
            blocking.append("Conflicting latest semester rows")
        if profile.get("Probation Status") == "Manual Review" and "Conflicting latest semester rows" not in blocking:
            blocking.append("Latest SGPA or LCGPA cannot be used for probation")
        if flags.get("conflicting_attempt"):
            blocking.append("Conflicting course attempt rows")
        if flags.get("withdrawal_conflict"):
            blocking.append("Withdrawal marker conflicts with a letter grade")
        if flags.get("result_conflict"):
            blocking.append("Passing letter conflicts with result F")
        if any("negative" in reason.casefold() for reason in reasons):
            blocking.append("A credit value is negative")
        if "Official credits attempted are missing or invalid" in reasons:
            blocking.append("Official credits attempted are missing or invalid")
        if "Multiple pre-holdout Student Summary rows" in reasons:
            blocking.append("Multiple pre-holdout student summary rows")
        if blocking:
            status = "Manual Review Required"
        elif information or review:
            status = "Valid with Missing Information"
        else:
            status = "Valid"
        rows.append({
            "Student Code": student_code,
            "Profile Status": status,
            "Notes": "; ".join(information + review + blocking),
        })
    return pd.DataFrame(rows)


def _course_note_flags(course_status: pd.DataFrame) -> dict[object, dict[str, bool]]:
    flags: dict[object, dict[str, bool]] = {}
    for status in course_status.to_dict(orient="records"):
        notes = set(_reason_parts(status.get("Review Notes")))
        student_flags = flags.setdefault(status["Student Code"], {})
        if "No passing-grade rule" in notes:
            student_flags["missing_rule"] = True
        if "Attempt outcome requires manual review" in notes:
            student_flags["unclassified"] = True
        if "Conflicting rows share a semester and attempt number" in notes:
            student_flags["conflicting_attempt"] = True
        if "Withdrawal marker conflicts with a letter grade" in notes:
            student_flags["withdrawal_conflict"] = True
        if "Passing letter conflicts with result F" in notes:
            student_flags["result_conflict"] = True
    return flags


def _presentation_issues(profiles: pd.DataFrame, course_status: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for profile in profiles.to_dict(orient="records"):
        student_code = profile["Student Code"]
        reasons = set(_reason_parts(profile.get("Manual Review Reasons")))
        if profile.get("Current Specialization") == "Not available before Spring 2026":
            rows.append(_issue(student_code, None, "Specialization", "Missing historical specialization", "Information"))
        if profile.get("Current Level") == "Unknown/Manual Review":
            rows.append(_issue(student_code, None, "Level", "Current level missing on latest historical semester", "Information"))
        if _missing(profile.get("Remaining Courses Count")):
            rows.append(_issue(
                student_code,
                None,
                "Remaining courses",
                "Study-plan pathway is not resolved, so the remaining-course count was not guessed",
                "Review",
            ))
        if "Specialization text is not in the known label list" in reasons:
            rows.append(_issue(student_code, None, "Specialization", "Specialization text needs review", "Information"))
        if profile.get("Current Specialization") == "Manual Review" or "Transcript specialization conflicts with the inferred specialization path" in reasons:
            rows.append(_issue(student_code, None, "Specialization", "Conflicting specialization values", "Blocking"))
        if "Latest semester has conflicting summary rows" in reasons or "Probation cannot be decided from conflicting semester rows" in reasons:
            rows.append(_issue(student_code, None, "Semester summary", "Conflicting latest semester rows", "Blocking"))
        if profile.get("Probation Status") == "Manual Review" and "Probation cannot be decided from conflicting semester rows" not in reasons:
            rows.append(_issue(student_code, None, "Probation", "Latest SGPA or LCGPA cannot be used for probation", "Blocking"))
        if any("negative" in reason.casefold() for reason in reasons):
            rows.append(_issue(student_code, None, "Credits", "A credit value is negative", "Blocking"))
        if "Official credits attempted are missing or invalid" in reasons:
            rows.append(_issue(student_code, None, "Credits", "Official credits attempted are missing or invalid", "Blocking"))
        if profile.get("Earned Credits Validation Status") == EARNED_MISMATCH:
            rows.append(_issue(student_code, None, "Credits", "Derived earned credits do not match the official historical total", "Review"))
        if profile.get("Earned Credits Validation Status") == EARNED_INSUFFICIENT:
            rows.append(_issue(student_code, None, "Credits", "Semester earned credits are insufficient for a historical total", "Review"))
        for credit_note in _reason_parts(profile.get("Credit Review Note")):
            rows.append(_issue(student_code, None, "Credits", credit_note, "Review"))
    for status in course_status.to_dict(orient="records"):
        notes = set(_reason_parts(status.get("Review Notes")))
        student_code = status["Student Code"]
        course_code = status["Course Code"]
        if "No passing-grade rule" in notes:
            rows.append(_issue(student_code, course_code, "Passing grade", "Missing passing-grade rule", "Review"))
        if "Attempt outcome requires manual review" in notes:
            rows.append(_issue(student_code, course_code, "Course outcome", "Course outcome could not be classified", "Review"))
        if "Conflicting rows share a semester and attempt number" in notes:
            rows.append(_issue(student_code, course_code, "Course history", "Conflicting attempt rows", "Blocking"))
        if "Withdrawal marker conflicts with a letter grade" in notes:
            rows.append(_issue(student_code, course_code, "Course history", "Withdrawal marker conflicts with a letter grade", "Blocking"))
        if "Passing letter conflicts with result F" in notes:
            rows.append(_issue(student_code, course_code, "Course history", "Passing letter conflicts with result F", "Blocking"))
    issues = pd.DataFrame(rows, columns=["Student Code", "Course Code", "Issue Type", "Issue", "Severity"])
    if issues.empty:
        return issues
    return issues.sort_values(["Student Code", "Course Code", "Issue"], kind="mergesort").reset_index(drop=True)


def _issue(student_code: object, course_code: object, issue_type: str, issue: str, severity: str) -> dict[str, object]:
    return {
        "Student Code": student_code,
        "Course Code": course_code,
        "Issue Type": issue_type,
        "Issue": issue,
        "Severity": severity,
    }


def _short_course_note(value: object) -> object:
    messages = []
    for part in _reason_parts(value):
        if part.startswith("Warning:"):
            continue
        if part == "No passing-grade rule":
            messages.append("No passing-grade rule for this course")
        elif part == "Attempt outcome requires manual review":
            messages.append("Course outcome could not be classified")
        elif part == "Conflicting rows share a semester and attempt number":
            messages.append("Conflicting attempt rows")
        elif part == "Withdrawal marker conflicts with a letter grade":
            messages.append("Withdrawal marker conflicts with a letter grade")
        elif part == "Passing letter conflicts with result F":
            messages.append("Passing letter conflicts with result F")
        else:
            messages.append(part)
    if not messages:
        return pd.NA
    return "; ".join(dict.fromkeys(messages))


def _reason_parts(value: object) -> list[str]:
    if _missing(value):
        return []
    return [part.strip() for part in str(value).split("|") if part.strip() and part.strip().casefold() != "nan"]


def _yes_no(values: pd.Series) -> pd.Series:
    return values.map(lambda value: "Yes" if bool(value) else "No")


def _write_validation_issues(worksheet, issues: pd.DataFrame) -> None:
    headers = ["Student Code", "Course Code", "Issue Type", "Issue", "Severity"]
    for column, header in enumerate(headers, start=4):
        worksheet.cell(1, column, header)
    for row_number, issue in enumerate(issues.itertuples(index=False), start=2):
        for column, value in enumerate(issue, start=4):
            worksheet.cell(row_number, column, None if _missing(value) else value)


def _write_quality_checks(worksheet, checks: pd.DataFrame) -> None:
    headers = ["Check", "Result", "Status"]
    for column, header in enumerate(headers, start=10):
        worksheet.cell(1, column, header)
    for row_number, check in enumerate(checks.itertuples(index=False), start=2):
        for column, value in enumerate(check, start=10):
            worksheet.cell(row_number, column, None if _missing(value) else value)


def _format_credit_audit_sheet(worksheet) -> None:
    widths = {
        "A": 16, "B": 32, "C": 40, "D": 16, "E": 42,
        "F": 26, "G": 42, "H": 62, "I": 62, "J": 62,
    }
    _format_sheet(
        worksheet,
        freeze="B2",
        widths=widths,
        text_columns={"A"},
        integer_columns=set(),
        decimal_columns={"B", "C", "D", "F", "G"},
        wrap_columns={"E", "H", "I", "J"},
    )


def _format_profile_sheet(worksheet) -> None:
    widths = {
        "A": 16, "B": 24, "C": 42, "D": 18, "E": 14, "F": 14, "G": 14,
        "H": 24, "I": 22, "J": 42, "K": 36, "L": 36, "M": 36, "N": 26, "O": 20, "P": 24,
    }
    _format_sheet(
        worksheet,
        freeze="B2",
        widths=widths,
        text_columns={"A"},
        integer_columns={"N"},
        decimal_columns={"E", "F", "G", "H", "I"},
        wrap_columns={"C", "J", "K", "L", "M"},
    )


def _format_course_sheet(worksheet) -> None:
    widths = {
        "A": 16, "B": 16, "C": 36, "D": 16, "E": 22, "F": 22, "G": 14, "H": 14,
        "I": 24, "J": 14, "K": 16, "L": 18, "M": 12, "N": 22, "O": 22, "P": 20, "Q": 46,
    }
    _format_sheet(
        worksheet,
        freeze="B2",
        widths=widths,
        text_columns={"A", "B"},
        integer_columns={"D"},
        decimal_columns=set(),
        wrap_columns={"C", "Q"},
    )


def _format_validation_sheet(worksheet, metric_rows: int, issue_rows: int) -> None:
    from openpyxl.styles import Alignment, Border, Font, Side

    header_font = Font(bold=True)
    header_border = Border(bottom=Side(style="thin", color="666666"))
    header_alignment = Alignment(vertical="center", wrap_text=True)
    for column in range(1, 13):
        cell = worksheet.cell(1, column)
        if cell.value:
            cell.font = header_font
            cell.border = header_border
            cell.alignment = header_alignment
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"D1:H{max(issue_rows + 1, 1)}"
    worksheet.row_dimensions[1].height = 22
    widths = {"A": 68, "B": 18, "C": 3, "D": 16, "E": 16, "F": 20, "G": 62, "H": 14, "I": 3, "J": 68, "K": 14, "L": 12}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    for row in worksheet.iter_rows(min_row=2, max_row=metric_rows + 1, min_col=2, max_col=2):
        value = row[0].value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            row[0].number_format = "0"
            row[0].alignment = Alignment(horizontal="right")
    for row in worksheet.iter_rows(min_row=2, max_row=max(issue_rows + 1, 2), min_col=4, max_col=5):
        for cell in row:
            if cell.value is not None:
                cell.number_format = "@"
    for row in worksheet.iter_rows(min_row=2, max_row=max(issue_rows + 1, 2), min_col=7, max_col=7):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _format_sheet(
    worksheet,
    *,
    freeze: str,
    widths: dict[str, int],
    text_columns: set[str],
    integer_columns: set[str],
    decimal_columns: set[str],
    wrap_columns: set[str],
) -> None:
    from openpyxl.styles import Alignment, Border, Font, Side

    header_font = Font(bold=True)
    header_border = Border(bottom=Side(style="thin", color="666666"))
    header_alignment = Alignment(vertical="center", wrap_text=True)
    for cell in worksheet[1]:
        cell.font = header_font
        cell.border = header_border
        cell.alignment = header_alignment
    worksheet.freeze_panes = freeze
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 22
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    last_row = worksheet.max_row
    for column in text_columns:
        for cell in worksheet[column][1:last_row]:
            if cell.value is not None:
                cell.number_format = "@"
                cell.value = str(cell.value)
    for column in integer_columns:
        for cell in worksheet[column][1:last_row]:
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = "0"
    for column in decimal_columns:
        for cell in worksheet[column][1:last_row]:
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = "0.00"
    wrap_alignment = Alignment(wrap_text=True, vertical="top")
    for column in wrap_columns:
        for cell in worksheet[column][1:last_row]:
            cell.alignment = wrap_alignment


def task7_timestamp() -> str:
    """Return a local timestamp for the run record."""
    return datetime.now().replace(microsecond=0).isoformat(sep=" ")


def _collapse_attempts(group: pd.DataFrame) -> tuple[list[dict[str, object]], list[str]]:
    ordered = group.sort_values(["Period Order", "Attempt Number", "Source Row ID"], kind="mergesort")
    notes: list[str] = []
    attempts: list[dict[str, object]] = []
    for (_, _), rows in ordered.groupby(["Period Order", "Attempt Number"], sort=False, dropna=False):
        conflict = _rows_conflict(rows)
        if len(rows) > 1 and conflict:
            notes.append("Conflicting rows share a semester and attempt number")
        elif len(rows) > 1:
            notes.append("Warning: identical duplicate attempt rows were collapsed")
        source = rows.iloc[0]
        grade = None if conflict else source["Grade Token"]
        result = None if conflict else source["Result Token"]
        grade_point = pd.NA if conflict else source["Grade Point"]
        withdrawal = bool(grade == "W" or result == "W")
        failed_letter = bool(grade in FAILURE_GRADES or result == "F")
        if result not in {None, "P", "F", "W"} and not conflict:
            notes.append("Warning: result text is not P, F, or W; the letter grade was used when it is on the scale")
        if conflict:
            meets = None
        elif withdrawal and grade in GRADE_RANK:
            meets = None
            notes.append("Withdrawal marker conflicts with a letter grade")
        else:
            meets = None
        attempts.append({
            "semester": source["Semester"],
            "grade": None if conflict else normalize_grade(source["Grade"]),
            "result": None if conflict else result,
            "grade_point": grade_point,
            "course_name": None if _missing(source["Course Name"]) else str(source["Course Name"]).strip(),
            "credit_hours": source["Credit Hours"] if "Credit Hours" in rows.columns else pd.NA,
            "withdrawal": withdrawal and not conflict,
            "failed_letter": failed_letter and not conflict,
            "meets": meets,
            "conflict": conflict,
            "period": int(source["Period Order"]),
        })
    return attempts, notes


def _summarize_group(group: pd.DataFrame, rule: Mapping[str, object] | None) -> dict[str, object]:
    student_code = group["Student Code"].iloc[0]
    course_code = str(group["Course Code"].iloc[0])
    attempts, notes = _collapse_attempts(group)
    minimum_grade = None if rule is None else normalize_grade(rule.get("Passing Grade"))
    minimum_point = None if rule is None else rule.get("Passing Grade Point")
    rule_resolved = minimum_grade in GRADE_RANK
    if not rule_resolved:
        notes.append("No passing-grade rule")
    for attempt in attempts:
        if attempt["conflict"] or not rule_resolved:
            attempt["meets"] = None
            continue
        if attempt["withdrawal"] and attempt["grade"] in GRADE_RANK:
            attempt["meets"] = None
            continue
        attempt["meets"] = meets_course_passing_requirement(
            course_code,
            attempt["grade"],
            attempt["grade_point"],
            {"Passing Grade": minimum_grade, "Passing Grade Point": minimum_point},
        )
        if attempt["meets"] is True and attempt["result"] == "F":
            attempt["meets"] = None
            attempt["conflict"] = True
            notes.append("Passing letter conflicts with result F")
    ever_failed = any(attempt["failed_letter"] for attempt in attempts)
    ever_withdrawn = any(attempt["withdrawal"] for attempt in attempts)
    first_pass = None
    if rule_resolved:
        for index, attempt in enumerate(attempts):
            if attempt["meets"] is True:
                first_pass = index
                break
    passed = first_pass is not None
    latest = attempts[-1]
    blocking_conflict = any(attempt["conflict"] for attempt in attempts) and not passed
    if not rule_resolved or blocking_conflict:
        final_status = "Manual Review"
    elif passed:
        final_status = "Completed"
    elif latest["withdrawal"]:
        final_status = "Withdrawn"
    elif latest["meets"] is False and latest["grade"] in FAILURE_GRADES:
        final_status = "Failed"
    elif latest["meets"] is False:
        final_status = "Below Passing Grade"
    else:
        final_status = "Manual Review"
    if passed and any(attempt["meets"] is None for attempt in attempts):
        notes.append("Warning: an unclassified attempt remains in the history; a valid pass determines completion")
    elif final_status == "Manual Review" and "No passing-grade rule" not in notes and "Attempt outcome requires manual review" not in notes:
        notes.append("Attempt outcome requires manual review")
    names = [attempt["course_name"] for attempt in attempts if attempt["course_name"]]
    credit_values = [attempt["credit_hours"] for attempt in attempts]
    credits_complete = all(classify_numeric_value(value, minimum=0, maximum=None) == "ok" for value in credit_values)
    completion = attempts[first_pass] if first_pass is not None else None
    completion_credit = completion["credit_hours"] if completion else pd.NA
    return {
        "Student Code": student_code,
        "Course Code": course_code,
        "Course Name": names[-1] if names else pd.NA,
        "Attempt Count": len(attempts),
        "First Attempt Semester": attempts[0]["semester"],
        "Latest Attempt Semester": latest["semester"],
        "First Grade": attempts[0]["grade"],
        "Latest Grade": latest["grade"],
        "First Result": attempts[0]["result"],
        "Latest Result": latest["result"],
        "Latest Grade Point": latest["grade_point"],
        "Minimum Passing Grade": minimum_grade if rule_resolved else pd.NA,
        "Minimum Passing Grade Point": minimum_point if rule_resolved else pd.NA,
        "Ever Passed": passed,
        "Ever Failed": ever_failed,
        "Ever Withdrawn": ever_withdrawn,
        "Eventually Passed": passed,
        "Final Status": final_status,
        "Is Completed": final_status == "Completed",
        "Is Current Failed": final_status in {"Failed", "Below Passing Grade"},
        "Is Current Withdrawn": final_status == "Withdrawn",
        "Is Repeated": len(attempts) > 1,
        "Completion Semester": completion["semester"] if completion else pd.NA,
        "Completion Grade": completion["grade"] if completion else pd.NA,
        "Completion Grade Point": completion["grade_point"] if completion else pd.NA,
        "Number of Attempts Before Completion": first_pass if completion else pd.NA,
        "Passed On First Attempt": bool(first_pass == 0) if completion else False,
        "Attempt Credit Hours": float(sum(float(value) for value in credit_values)) if credits_complete else pd.NA,
        "Completion Credit Hours": completion_credit if completion and classify_numeric_value(completion_credit, minimum=0, maximum=None) == "ok" else pd.NA,
        "Review Notes": _join_reasons(*notes),
    }


def _rows_conflict(rows: pd.DataFrame) -> bool:
    if len(rows) == 1:
        return False
    compared = ["Grade Token", "Result Token", "Grade Point", "Credit Hours"]
    existing = [column for column in compared if column in rows.columns]
    for column in existing:
        values = rows[column].map(_comparison_token)
        if values.nunique(dropna=False) > 1:
            return True
    return False


def _numeric_difference(official: object, reconstructed: object) -> object:
    """Return reconstructed minus official when both values are finite numbers."""
    if classify_numeric_value(official, minimum=None, maximum=None) != "ok":
        return pd.NA
    if classify_numeric_value(reconstructed, minimum=None, maximum=None) != "ok":
        return pd.NA
    return float(reconstructed) - float(official)


def _comparison_token(value: object) -> str:
    if _missing(value):
        return ""
    return str(value).strip()


def _aggregate_course_lists(course_status: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for student_code, group in course_status.groupby("Student Code", sort=True):
        completed_count, completed_list = _code_list(group, "Is Completed")
        failed_count, failed_list = _code_list(group, "Is Current Failed")
        withdrawn_count, withdrawn_list = _code_list(group, "Is Current Withdrawn")
        repeated_count, repeated_list = _code_list(group, "Is Repeated")
        attempted_credits = _sum_or_blank(group["Attempt Credit Hours"])
        earned_credits = _sum_or_blank(group.loc[group["Is Completed"], "Completion Credit Hours"])
        difference_attempted = pd.NA
        difference_earned = pd.NA
        course_notes = sorted({
            note.removeprefix("Warning: ").strip()
            for notes in group["Review Notes"].fillna("")
            for note in str(notes).split("|")
            if note.strip() and not note.strip().startswith("Warning:")
        })
        rows.append({
            "Student Code": student_code,
            "Completed Course Count": completed_count,
            "Completed Courses List": completed_list,
            "Failed Course Count": failed_count,
            "Failed Courses List": failed_list,
            "Withdrawn Course Count": withdrawn_count,
            "Withdrawn Courses List": withdrawn_list,
            "Repeated Course Count": repeated_count,
            "Repeated Courses List": repeated_list,
            "Historical Attempt Count": int(group["Attempt Count"].sum()),
            "Unique Courses Attempted": int(group["Course Code"].nunique()),
            "Reconstructed Credits Attempted": attempted_credits,
            "Reconstructed Credits Earned": earned_credits,
            "Attempted Credits Difference": difference_attempted,
            "Earned Credits Difference": difference_earned,
            "Course Review Reasons": _join_reasons(*course_notes),
        })
    return pd.DataFrame(rows)


def _profile_review(row: pd.Series) -> pd.Series:
    reasons: list[str] = []
    if bool(row.get("Snapshot Conflict")):
        reasons.append("Latest semester has conflicting summary rows")
    if row.get("Current Level") == "Unknown/Manual Review":
        reasons.append("Current level is missing on the latest historical semester")
    attribute_reason = row.get("Attribute Review Reason")
    if not _missing(attribute_reason):
        reasons.append(str(attribute_reason))
    if row.get("Probation Status") == "Manual Review" and not _missing(row.get("Probation Review Reason")):
        reasons.append(str(row["Probation Review Reason"]))
    for column, label in (
        ("Official Credits Attempted", "Official credits attempted"),
        ("Reconstructed Credits Attempted", "Reconstructed credits attempted"),
        ("Reconstructed Credits Earned", "Reconstructed credits earned"),
    ):
        status = classify_numeric_value(row.get(column), minimum=0, maximum=None)
        if status == "negative":
            reasons.append(f"{label} are negative")
        elif column == "Official Credits Attempted" and status != "ok":
            reasons.append("Official credits attempted are missing or invalid")
    course_reasons = row.get("Course Review Reasons")
    if not _missing(course_reasons):
        reasons.extend([part.strip() for part in str(course_reasons).split("|") if part.strip()])
    unique_reasons = list(dict.fromkeys(reasons))
    required = "Yes" if unique_reasons else "No"
    status = "Manual Review" if unique_reasons else "OK"
    return pd.Series([status, required, _join_reasons(*unique_reasons)])


def _probation_from_row(row: pd.Series) -> pd.Series:
    if bool(row.get("Snapshot Conflict")):
        return pd.Series(["Manual Review", pd.NA, "Probation cannot be decided from conflicting semester rows"])
    sgpa_status = classify_numeric_value(row.get("Latest SGPA"), minimum=0, maximum=4)
    lcgpa_status = classify_numeric_value(row.get("Latest LCGPA"), minimum=0, maximum=4)
    if sgpa_status != "ok" or lcgpa_status != "ok":
        return pd.Series(["Manual Review", pd.NA, "Latest SGPA or LCGPA is missing or outside 0 to 4"])
    sgpa = float(row["Latest SGPA"])
    lcgpa = float(row["Latest LCGPA"])
    if sgpa < PROBATION_THRESHOLD or lcgpa < PROBATION_THRESHOLD:
        return pd.Series(["Probation", True, pd.NA])
    return pd.Series(["Not Probation", False, pd.NA])


def _risk_category(probation_status: object) -> str:
    """Map the historical probation result to a research label.

    Academic Risk Category is a research-derived label based only on the
    historical probation rule and is not an official university risk classification.
    """
    if probation_status == "Probation":
        return "High Academic Risk"
    if probation_status == "Not Probation":
        return "No Probation Flag"
    return "Manual Review"


def _level_label(value: object) -> str:
    if _missing(value) or str(value).strip() == "":
        return "Unknown/Manual Review"
    return str(value).strip()


def _agreed_specialization(transcript_value: object, inferred_value: object) -> tuple[object, object]:
    transcript_label, transcript_reason = _canonical_specialization(transcript_value)
    inferred_label, inferred_reason = _canonical_specialization(inferred_value)
    if transcript_label is None and inferred_label is None:
        return pd.NA, "Specialization is missing on the pre-holdout student summary"
    if transcript_label is not None and inferred_label is not None and transcript_label != inferred_label:
        return pd.NA, "Transcript specialization conflicts with the inferred specialization path"
    label = transcript_label if transcript_label is not None else inferred_label
    reason = transcript_reason or inferred_reason
    return label, reason


def _canonical_specialization(value: object) -> tuple[str | None, str | None]:
    if _missing(value) or str(value).strip() == "":
        return None, None
    key = re.sub(r"\s+", " ", str(value).strip().casefold())
    if key in SPECIALIZATION_LABELS:
        return SPECIALIZATION_LABELS[key], None
    return str(value).strip(), "Specialization text is not in the known label list"


def _code_list(group: pd.DataFrame, flag: str) -> tuple[int, str]:
    codes = sorted(set(group.loc[group[flag].fillna(False).astype(bool), "Course Code"].astype(str)))
    if any(LIST_SEPARATOR.strip() in code or ";" in code for code in codes):
        raise RuntimeError("A course code contains the list separator.")
    return len(codes), LIST_SEPARATOR.join(codes)


def _assert_list_counts(profiles: pd.DataFrame, count_column: str, list_column: str) -> None:
    for record in profiles[[count_column, list_column, "Student Code"]].itertuples(index=False):
        count = int(record[0])
        text = "" if _missing(record[1]) else str(record[1]).strip()
        parts = [] if text == "" else [part.strip() for part in text.split(";")]
        if any(part == "" for part in parts) or len(parts) != len(set(parts)) or parts != sorted(parts):
            raise RuntimeError(f"{list_column} is not a unique sorted list for {record[2]}.")
        if count != len(parts):
            raise RuntimeError(f"{count_column} does not match {list_column} for {record[2]}.")


def _assert_non_negative(frame: pd.DataFrame, column: str) -> None:
    for value in frame[column]:
        status = classify_numeric_value(value, minimum=0, maximum=None)
        if status == "negative":
            raise RuntimeError(f"{column} contains a negative value.")


def _sum_or_blank(values: pd.Series) -> object:
    if values.isna().any():
        return pd.NA
    if len(values) == 0:
        return 0.0
    total = float(pd.to_numeric(values, errors="coerce").sum())
    return total


def _subset(frame: pd.DataFrame, flag: str, columns: list[str]) -> pd.DataFrame:
    selected = frame.loc[frame[flag].fillna(False).astype(bool), columns].copy()
    return selected.sort_values(["Student Code", "Course Code"], kind="mergesort").reset_index(drop=True)


def _order_columns(frame: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    missing = [column for column in columns if column not in frame.columns]
    if missing:
        raise RuntimeError("Required output columns are missing: " + ", ".join(missing))
    ordered = frame.loc[:, columns].copy()
    return ordered.sort_values(["Student Code"] + (["Course Code"] if "Course Code" in ordered.columns else []), kind="mergesort").reset_index(drop=True)


def _passing_rule_map(passing_rules: pd.DataFrame) -> dict[str, dict[str, object]]:
    work = passing_rules.copy()
    work["Course Code"] = work["Course Code"].map(normalize_course_code)
    work = work.dropna(subset=["Course Code"])
    rules: dict[str, dict[str, object]] = {}
    conflicts: list[str] = []
    for course_code, group in work.groupby("Course Code", sort=True):
        grades = {grade for grade in group["Passing Grade"].map(normalize_grade) if grade is not None}
        points = {point for point in group["Passing Grade Point"].map(_point_token)}
        if len(grades) != 1 or len(points) != 1:
            conflicts.append(str(course_code))
            continue
        grade = next(iter(grades))
        point = next(iter(points))
        rules[str(course_code)] = {
            "Passing Grade": grade,
            "Passing Grade Point": point,
        }
    if conflicts:
        raise RuntimeError("Passing-grade rules disagree for: " + ", ".join(sorted(conflicts)))
    return rules


def _point_token(value: object) -> float | None:
    if classify_numeric_value(value, minimum=None, maximum=None) != "ok":
        return None
    return float(value)


def _reject_holdout(frame: pd.DataFrame, label: str) -> pd.DataFrame:
    classified = classify_period_frame(frame)
    leakage = ~classified["Partition"].eq("Before Spring 2026")
    if leakage.any():
        sample = classified.loc[leakage, ["Student Code", "Semester", "Partition"]].head(5)
        raise RuntimeError(
            f"Spring 2026 or unresolved rows entered the Task 7 {label} pipeline:\n{sample.to_string(index=False)}"
        )
    if classified["Semester"].astype(str).eq(HOLDOUT_SEMESTER).any() or classified["Semester"].astype(str).str.contains("2026").any():
        raise RuntimeError(f"{HOLDOUT_SEMESTER} or another 2026 semester entered the Task 7 {label} pipeline.")
    return classified


def _boundary_semester(frame: pd.DataFrame, *, minimum: bool) -> str:
    order = frame["Period Order"].min() if minimum else frame["Period Order"].max()
    labels = sorted(set(frame.loc[frame["Period Order"].eq(order), "Semester"].astype(str)))
    if len(labels) != 1:
        raise RuntimeError(f"Boundary period {order} has more than one semester label: {labels}")
    return labels[0]


def _split_semester_text(value: object) -> tuple[int, str] | None:
    if _missing(value):
        return None
    parts = str(value).strip().split()
    if len(parts) != 2 or not parts[0].isdigit():
        return None
    if parts[1].casefold() not in {"spring", "fall"}:
        return None
    return int(parts[0]), parts[1]


def _result_token(value: object) -> str | None:
    token = normalize_grade(value)
    if token in {"P", "F", "W"}:
        return token
    return None if token is None else token


def _require_columns(frame: pd.DataFrame, columns: list[str], label: str) -> None:
    missing = [column for column in columns if column not in frame.columns]
    if missing:
        raise RuntimeError(f"{label} is missing columns: " + ", ".join(missing))


def _join_reasons(*reasons: object) -> object:
    cleaned = []
    for reason in reasons:
        if _missing(reason):
            continue
        text = str(reason).strip()
        if text and text not in cleaned:
            cleaned.append(text)
    if not cleaned:
        return pd.NA
    return " | ".join(cleaned)


def _missing(value: object) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False
